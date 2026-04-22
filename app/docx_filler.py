"""Fill a .docx template using an LLM and a JSON data source.

Two modes are supported:

* **Placeholder mode** (default when at least one ``{{name}}``, ``{name}`` or
  ``[name]`` marker is found in the document): the LLM is asked for a
  ``{placeholder: value}`` mapping, then placeholders are substituted in place.
* **Free-form mode** (when no placeholders are detected): the document is
  serialised as a numbered list of paragraphs and the LLM is asked to return
  edited versions of the paragraphs that should be filled in, using the JSON
  data as ground truth. Paragraphs the LLM does not touch are kept intact.

The module only deals with text content (paragraphs + table cell paragraphs);
images, headers, footers, and other structures are left untouched.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from .gemini_client import LLMClient

# Matches {{name}}, {name} (single braces), or [name]. ``name`` is captured
# without the surrounding delimiters.
_PLACEHOLDER_RE = re.compile(
    r"""
    (?P<full>
        \{\{\s*(?P<double>[^{}]+?)\s*\}\}      # {{ name }}
        |
        (?<!\{)\{\s*(?P<single>[^{}\n]+?)\s*\}(?!\})   # { name } (not {{ }})
        |
        \[\s*(?P<bracket>[^\[\]\n]+?)\s*\]     # [ name ]
    )
    """,
    re.VERBOSE,
)


@dataclass
class FillResult:
    """Outcome of a document fill operation."""

    content: bytes
    mode: str  # "placeholder" or "freeform"
    filled_placeholders: dict[str, str]
    changed_paragraphs: int


def _iter_paragraphs(doc: DocxDocument):
    """Yield every ``Paragraph`` in the document, including inside tables."""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)


def _iter_cell_paragraphs(cell: _Cell):
    yield from cell.paragraphs
    for table in cell.tables:
        for row in table.rows:
            for sub_cell in row.cells:
                yield from _iter_cell_paragraphs(sub_cell)


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) or paragraph.text


def _set_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    """Replace a paragraph's text while keeping the first run's formatting.

    All runs are collapsed into the first run; any extras are cleared. This
    loses per-run formatting (bold/italic inside a paragraph) but keeps the
    paragraph-level style, which is what users expect when filling templates.
    """
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for extra in runs[1:]:
        extra.text = ""


def find_placeholders(doc: DocxDocument) -> list[str]:
    """Return the list of unique placeholder names found in the document.

    Order is preserved (first occurrence wins).
    """
    seen: dict[str, None] = {}
    for paragraph in _iter_paragraphs(doc):
        text = _paragraph_text(paragraph)
        for match in _PLACEHOLDER_RE.finditer(text):
            name = match.group("double") or match.group("single") or match.group("bracket")
            if name is None:
                continue
            name = name.strip()
            if name and name not in seen:
                seen[name] = None
    return list(seen.keys())


def _replace_placeholders_in_paragraph(
    paragraph: Paragraph, values: dict[str, str]
) -> bool:
    """Replace every placeholder in ``paragraph`` with the matching value.

    Returns True if the paragraph was changed.
    """
    original = _paragraph_text(paragraph)
    if not original:
        return False

    def _substitute(match: re.Match[str]) -> str:
        name = (
            match.group("double")
            or match.group("single")
            or match.group("bracket")
            or ""
        ).strip()
        if name in values:
            return str(values[name])
        return match.group("full")

    replaced = _PLACEHOLDER_RE.sub(_substitute, original)
    if replaced == original:
        return False
    _set_paragraph_text(paragraph, replaced)
    return True


def _apply_placeholder_values(doc: DocxDocument, values: dict[str, str]) -> int:
    changed = 0
    for paragraph in _iter_paragraphs(doc):
        if _replace_placeholders_in_paragraph(paragraph, values):
            changed += 1
    return changed


_PLACEHOLDER_SYSTEM_PROMPT = """You are an assistant that fills Microsoft Word document templates.

You will be given:
1. A list of placeholder names extracted from a .docx template.
2. A JSON payload that may be either a direct mapping of values, a knowledge base the assistant must mine for values, or a mix of both.

Your job: return a JSON object mapping every placeholder name to the best value
for that placeholder, inferred from the JSON payload. If a placeholder cannot
be resolved from the payload, return an empty string for it. Values must be
plain strings (dates, numbers, etc. formatted as the user would expect to read
them in a finished document). Do not invent facts that are not present in the
payload. Preserve the original language of the document.

Respond with a single JSON object whose keys are the placeholder names and
whose values are the resolved strings. No prose."""


_FREEFORM_SYSTEM_PROMPT = """You are an assistant that fills Microsoft Word documents.

You will be given:
1. The document as an ordered list of paragraphs, each with an integer id.
2. A JSON payload that may be either a direct mapping of values, a knowledge base the assistant must mine for values, or a mix of both.

Your job: identify paragraphs that contain blank spots, fill-in prompts, underscores, or questions that should be answered using the JSON payload, and return edited versions of ONLY those paragraphs.

Respond with a single JSON object of the form:
{"edits": [{"id": <int>, "new_text": "..."}, ...]}

Rules:
- Only include paragraphs you actually modify. Leave untouched ones out.
- Preserve the original language of the document.
- Do not invent facts that are not present in the payload. If a field cannot be filled, leave the paragraph out.
- Keep the overall structure and tone of the document."""


def _run_placeholder_mode(
    doc: DocxDocument,
    placeholders: list[str],
    data: Any,
    llm: LLMClient,
) -> tuple[dict[str, str], int]:
    user_prompt = json.dumps(
        {
            "placeholders": placeholders,
            "data": data,
        },
        ensure_ascii=False,
        indent=2,
    )
    raw = llm.complete_json(_PLACEHOLDER_SYSTEM_PROMPT, user_prompt)
    values: dict[str, str] = {}
    for name in placeholders:
        v = raw.get(name, "")
        values[name] = "" if v is None else str(v)
    changed = _apply_placeholder_values(doc, values)
    return values, changed


def _run_freeform_mode(
    doc: DocxDocument,
    data: Any,
    llm: LLMClient,
) -> int:
    paragraphs = list(_iter_paragraphs(doc))
    numbered = [
        {"id": idx, "text": _paragraph_text(p)}
        for idx, p in enumerate(paragraphs)
        if _paragraph_text(p).strip()
    ]
    if not numbered:
        return 0
    user_prompt = json.dumps(
        {"paragraphs": numbered, "data": data},
        ensure_ascii=False,
        indent=2,
    )
    raw = llm.complete_json(_FREEFORM_SYSTEM_PROMPT, user_prompt)
    edits = raw.get("edits", [])
    if not isinstance(edits, list):
        raise RuntimeError("LLM free-form response must have 'edits' as a list.")
    changed = 0
    for edit in edits:
        if not isinstance(edit, dict):
            continue
        idx = edit.get("id")
        new_text = edit.get("new_text")
        if not isinstance(idx, int) or not isinstance(new_text, str):
            continue
        if 0 <= idx < len(paragraphs):
            _set_paragraph_text(paragraphs[idx], new_text)
            changed += 1
    return changed


def fill_document(docx_bytes: bytes, data: Any, llm: LLMClient) -> FillResult:
    """Fill a .docx document using the provided JSON ``data`` and ``llm``.

    ``data`` is typically the parsed JSON file sent by the caller and may be
    any JSON-serialisable value.
    """
    doc = Document(io.BytesIO(docx_bytes))
    placeholders = find_placeholders(doc)

    filled: dict[str, str] = {}
    if placeholders:
        filled, changed = _run_placeholder_mode(doc, placeholders, data, llm)
        mode = "placeholder"
    else:
        changed = _run_freeform_mode(doc, data, llm)
        mode = "freeform"

    buffer = io.BytesIO()
    doc.save(buffer)
    return FillResult(
        content=buffer.getvalue(),
        mode=mode,
        filled_placeholders=filled,
        changed_paragraphs=changed,
    )

"""Generate the example .docx template used in the README.

Run from the repo root: `python examples/make_example.py`.
"""

from pathlib import Path

from docx import Document


def main() -> None:
    out_dir = Path(__file__).parent
    doc = Document()
    doc.add_heading("Коммерческое предложение", level=1)
    doc.add_paragraph("Уважаемый(ая) {{ client_name }}!")
    doc.add_paragraph(
        "Благодарим Вас за интерес к компании {{ company_name }}. "
        "Рады предложить услугу «{{ service }}» на следующих условиях:"
    )
    doc.add_paragraph("Стоимость: { amount } {{ currency }}")
    doc.add_paragraph("Срок выполнения: [ deadline ]")
    doc.add_paragraph("Контактное лицо: {{ contact_person }}, {{ contact_email }}")
    doc.add_paragraph()
    doc.add_paragraph("С уважением,\nКоманда {{ company_name }}")
    doc.save(out_dir / "example_template.docx")
    print(f"Wrote {out_dir / 'example_template.docx'}")


if __name__ == "__main__":
    main()

"""Tests for :mod:`app.data_sources`."""

from __future__ import annotations

import io
import json

import pytest
from docx import Document
from openpyxl import Workbook

from app.data_sources import UnsupportedDataFormat, parse_data_file


def test_parse_json_utf8() -> None:
    payload = {"имя": "Иван", "сумма": 100}
    content = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    assert parse_data_file("data.json", content) == payload


def test_parse_json_with_bom() -> None:
    payload = {"a": 1}
    content = "\ufeff".encode() + json.dumps(payload).encode("utf-8")
    assert parse_data_file("data.json", content) == payload


def test_parse_json_invalid_raises() -> None:
    with pytest.raises(UnsupportedDataFormat):
        parse_data_file("bad.json", b"not json at all")


def test_parse_text_plain() -> None:
    content = b"Hello\nWorld"
    result = parse_data_file("notes.txt", content)
    assert result == {"text": "Hello\nWorld"}


def test_parse_text_cp1251() -> None:
    content = "Привет мир".encode("cp1251")
    result = parse_data_file("notes.txt", content)
    assert result == {"text": "Привет мир"}


def test_parse_markdown() -> None:
    content = b"# Title\n\nBody"
    assert parse_data_file("doc.md", content) == {"text": "# Title\n\nBody"}


def test_parse_csv_with_header() -> None:
    content = b"name,age\nAlice,30\nBob,25\n"
    result = parse_data_file("people.csv", content)
    assert result == {"rows": [{"name": "Alice", "age": "30"}, {"name": "Bob", "age": "25"}]}


def test_parse_csv_without_header_falls_back_to_list_rows() -> None:
    # Duplicate header values -> not treated as header.
    content = b"a,a\n1,2\n3,4\n"
    result = parse_data_file("raw.csv", content)
    assert result == {"rows": [["a", "a"], ["1", "2"], ["3", "4"]]}


def test_parse_csv_semicolon_delimiter() -> None:
    content = "имя;значение\nИван;100\nМаша;200\n".encode()
    result = parse_data_file("data.csv", content)
    assert result == {
        "rows": [{"имя": "Иван", "значение": "100"}, {"имя": "Маша", "значение": "200"}]
    }


def test_parse_xlsx_multi_sheet() -> None:
    wb = Workbook()
    sheet1 = wb.active
    sheet1.title = "People"
    sheet1.append(["name", "age"])
    sheet1.append(["Alice", 30])
    sheet1.append(["Bob", 25])
    sheet2 = wb.create_sheet("Notes")
    sheet2.append(["hello"])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_data_file("book.xlsx", buf.getvalue())
    assert result == {
        "sheets": {
            "People": [["name", "age"], ["Alice", 30], ["Bob", 25]],
            "Notes": [["hello"]],
        }
    }


def test_parse_xlsx_skips_blank_rows() -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "S"
    sheet.append(["a", "b"])
    sheet.append([None, None])
    sheet.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_data_file("book.xlsx", buf.getvalue())
    assert result == {"sheets": {"S": [["a", "b"], [1, 2]]}}


def test_parse_docx_paragraphs_and_tables() -> None:
    doc = Document()
    doc.add_paragraph("Первый абзац")
    doc.add_paragraph("")  # blank, should be dropped
    doc.add_paragraph("Второй абзац")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Имя"
    table.cell(0, 1).text = "Иван"
    table.cell(1, 0).text = "Возраст"
    table.cell(1, 1).text = "42"
    buf = io.BytesIO()
    doc.save(buf)

    result = parse_data_file("knowledge.docx", buf.getvalue())
    assert result == {
        "paragraphs": ["Первый абзац", "Второй абзац"],
        "tables": [[["Имя", "Иван"], ["Возраст", "42"]]],
    }


def test_unsupported_extension() -> None:
    with pytest.raises(UnsupportedDataFormat):
        parse_data_file("data.xml", b"<root/>")


def test_empty_filename() -> None:
    with pytest.raises(UnsupportedDataFormat):
        parse_data_file("", b"content")
